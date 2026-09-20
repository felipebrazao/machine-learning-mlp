from pathlib import Path

from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "apresentacao"
FIG_DIR = ROOT / "notebooks" / "figuras"
OUT_FILE = OUT_DIR / "Apresentacao_MLP_Veiculos.pptx"

W = Inches(13.333)
H = Inches(7.5)

NAVY = RGBColor(15, 23, 42)
BLUE = RGBColor(37, 99, 235)
CYAN = RGBColor(14, 165, 233)
TEAL = RGBColor(13, 148, 136)
GREEN = RGBColor(22, 163, 74)
AMBER = RGBColor(245, 158, 11)
RED = RGBColor(220, 38, 38)
WHITE = RGBColor(248, 250, 252)
MUTED = RGBColor(148, 163, 184)
PANEL = RGBColor(30, 41, 59)
PANEL_2 = RGBColor(51, 65, 85)


def rect(slide, x, y, w, h, fill, radius=True, line=None):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    return shape


def text(slide, value, x, y, w, h, size=20, color=WHITE, bold=False,
         align=PP_ALIGN.LEFT, font="Aptos", valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def rich_lines(slide, lines, x, y, w, h, size=18, gap=7, bullet=True):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    for i, line in enumerate(lines):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = line
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = WHITE
        p.space_after = Pt(gap)
        p.level = 0
        if bullet:
            p.text = f"•  {line}"
    return box


def base_slide(prs, title, kicker=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = NAVY
    rect(slide, 0, 0, 0.16, 7.5, BLUE, radius=False)
    if kicker:
        text(slide, kicker.upper(), 0.62, 0.34, 4.5, 0.28, 10, CYAN, True)
    text(slide, title, 0.62, 0.63, 12.0, 0.62, 29, WHITE, True)
    return slide


def footer(slide, number):
    text(slide, "Regressão de preços de veículos com MLP", 0.62, 7.12, 6.2, 0.2, 9, MUTED)
    text(slide, f"{number:02d}", 12.1, 7.08, 0.55, 0.24, 10, MUTED, True, PP_ALIGN.RIGHT)


def metric_card(slide, label, value, x, y, w=2.45, accent=BLUE, note=None):
    rect(slide, x, y, w, 1.25, PANEL)
    rect(slide, x, y, 0.08, 1.25, accent, radius=False)
    text(slide, label.upper(), x + 0.22, y + 0.18, w - 0.38, 0.22, 9, MUTED, True)
    text(slide, value, x + 0.22, y + 0.46, w - 0.38, 0.42, 22, WHITE, True)
    if note:
        text(slide, note, x + 0.22, y + 0.94, w - 0.38, 0.18, 8, MUTED)


def add_picture_contain(slide, path, x, y, w, h):
    from PIL import Image

    with Image.open(path) as img:
        iw, ih = img.size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    slide.shapes.add_picture(str(path), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2),
                             width=Inches(pw), height=Inches(ph))


def add_bar_chart(slide, categories, series, x, y, w, h, max_scale=None):
    rect(slide, x, y, w, h, WHITE)
    data = ChartData()
    data.categories = categories
    for name, values in series:
        data.add_series(name, values)
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), data
    ).chart
    chart.has_legend = len(series) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.font.size = Pt(10)
        chart.legend.font.color.rgb = MUTED
    chart.has_title = False
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.major_gridlines.format.line.color.rgb = MUTED
    chart.value_axis.tick_labels.font.color.rgb = PANEL_2
    chart.value_axis.tick_labels.font.size = Pt(10)
    chart.category_axis.tick_labels.font.color.rgb = NAVY
    chart.category_axis.tick_labels.font.size = Pt(11)
    if max_scale:
        chart.value_axis.maximum_scale = max_scale
    palette = [BLUE, CYAN, TEAL]
    for i, ser in enumerate(chart.series):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = palette[i % len(palette)]
        ser.format.line.color.rgb = palette[i % len(palette)]
    return chart


def sensitivity_slide(prs, number, person, parameter, image, conclusion, values):
    slide = base_slide(prs, person, "Análise individual de sensibilidade")
    text(slide, parameter, 0.64, 1.26, 4.9, 0.34, 15, CYAN, True)
    rect(slide, 0.62, 1.72, 7.85, 4.95, WHITE)
    add_picture_contain(slide, image, 0.79, 1.89, 7.51, 4.6)
    rect(slide, 8.72, 1.72, 3.95, 2.15, PANEL)
    text(slide, "MELHOR PONTO", 9.02, 2.03, 3.2, 0.25, 11, MUTED, True)
    text(slide, values[0], 9.02, 2.40, 3.2, 0.48, 25, WHITE, True)
    text(slide, values[1], 9.02, 3.02, 3.2, 0.27, 12, GREEN, True)
    rect(slide, 8.72, 4.08, 3.95, 2.59, PANEL)
    text(slide, "LEITURA", 9.02, 4.38, 3.2, 0.25, 11, MUTED, True)
    text(slide, conclusion, 9.02, 4.78, 3.15, 1.45, 15, WHITE)
    footer(slide, number)


def main():
    OUT_DIR.mkdir(exist_ok=True)
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    # 1 — Capa
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    rect(slide, 0, 0, 0.18, 7.5, BLUE, radius=False)
    text(slide, "INTELIGÊNCIA COMPUTACIONAL · MLP", 0.72, 0.7, 6.5, 0.35, 12, CYAN, True)
    text(slide, "Regressão de preços\nde veículos com MLP", 0.72, 1.25, 7.5, 1.7, 35, WHITE, True)
    text(slide, "Comparação entre Random Search e TPE, com análise de sensibilidade individual", 0.74, 3.25, 7.0, 0.85, 18, MUTED)
    rect(slide, 8.55, 0.85, 3.7, 5.65, PANEL)
    text(slide, "558.837", 8.95, 1.42, 2.9, 0.62, 34, WHITE, True)
    text(slide, "registros", 8.97, 2.05, 2.6, 0.3, 13, MUTED)
    text(slide, "25 + 25", 8.95, 2.86, 2.9, 0.62, 34, CYAN, True)
    text(slide, "configurações avaliadas", 8.97, 3.5, 2.6, 0.3, 13, MUTED)
    text(slide, "R² 0,971", 8.95, 4.32, 2.9, 0.62, 34, GREEN, True)
    text(slide, "no teste reservado", 8.97, 4.95, 2.6, 0.3, 13, MUTED)
    text(slide, "Andrey · Felipe Gabriel · Felipe Vieira · Gabriel", 0.74, 6.67, 9.2, 0.3, 13, WHITE)
    footer(slide, 1)

    # 2 — Equipe e IA
    slide = base_slide(prs, "Equipe, escopo e uso de IA", "Transparência")
    names = [
        "Andrey de Matos Gonçalves",
        "Felipe Gabriel Souza Libório",
        "Felipe Vieira Brazão e Silva",
        "Gabriel Coelho Goes",
    ]
    for i, name in enumerate(names):
        y = 1.52 + i * 1.12
        rect(slide, 0.62, y, 5.72, 0.86, PANEL)
        text(slide, f"0{i + 1}", 0.88, y + 0.21, 0.52, 0.27, 12, CYAN, True)
        text(slide, name, 1.48, y + 0.16, 4.5, 0.36, 17, WHITE, True)
    rect(slide, 6.67, 1.52, 6.0, 4.22, PANEL)
    text(slide, "DECLARAÇÃO DE IA GENERATIVA", 7.03, 1.87, 5.25, 0.28, 12, CYAN, True)
    rich_lines(slide, [
        "Ferramenta: Codex, da OpenAI",
        "Apoio na estruturação, revisão, documentação e interpretação",
        "Participação aproximada: dois terços do desenvolvimento textual e técnico",
        "Resultados conferidos por execução, validação cruzada e revisão humana",
    ], 7.02, 2.35, 5.05, 2.85, 16, 10)
    text(slide, "Responsabilidade final e interpretação: equipe.", 7.03, 5.22, 5.0, 0.28, 12, AMBER, True)
    footer(slide, 2)

    # 3 — Dataset
    slide = base_slide(prs, "Problema e conjunto de dados", "Contexto")
    text(slide, "Objetivo", 0.65, 1.48, 2.5, 0.3, 13, CYAN, True)
    text(slide, "Estimar o preço efetivo de venda de um veículo a partir de características do automóvel, condição e informações de mercado.", 0.65, 1.88, 5.65, 1.25, 21, WHITE, True)
    metric_card(slide, "Instâncias", "558.837", 0.65, 3.52, 2.55, BLUE)
    metric_card(slide, "Colunas", "16", 3.42, 3.52, 2.55, CYAN, "15 preditores + alvo")
    metric_card(slide, "Tipo", "Regressão", 0.65, 5.02, 2.55, TEAL)
    metric_card(slide, "Alvo", "sellingprice", 3.42, 5.02, 2.55, AMBER, "valor em dólares")
    rect(slide, 6.65, 1.48, 6.02, 4.79, PANEL)
    text(slide, "VEHICLE SALES AND MARKET TRENDS", 7.05, 1.88, 5.2, 0.35, 14, CYAN, True)
    rich_lines(slide, [
        "Fonte pública: Kaggle",
        "Veículos vendidos em leilões nos Estados Unidos",
        "Atributos numéricos, categóricos e temporais",
        "Questão central: quão próximo o modelo chega do preço real?",
    ], 7.02, 2.43, 5.0, 2.65, 17, 11)
    text(slide, "kaggle.com/datasets/syedanwarafridi/vehicle-sales-data", 7.03, 5.51, 5.0, 0.35, 10, MUTED)
    footer(slide, 3)

    # 4 — Pipeline
    slide = base_slide(prs, "Pipeline experimental sem vazamento", "Metodologia")
    steps = [
        ("01", "Holdout", "20% reservado antes de qualquer ajuste"),
        ("02", "Preparação", "imputação, one-hot e padronização"),
        ("03", "Validação", "3 folds somente no treino"),
        ("04", "Seleção", "menor MAE médio em dólares"),
        ("05", "Teste final", "uma única avaliação no holdout"),
    ]
    for i, (num, title, desc) in enumerate(steps):
        x = 0.62 + i * 2.48
        rect(slide, x, 1.65, 2.18, 3.16, PANEL)
        text(slide, num, x + 0.24, 1.96, 0.55, 0.35, 13, CYAN, True)
        text(slide, title, x + 0.24, 2.51, 1.72, 0.5, 18, WHITE, True)
        text(slide, desc, x + 0.24, 3.25, 1.66, 1.05, 14, MUTED)
        if i < 4:
            text(slide, "→", x + 2.17, 2.88, 0.32, 0.4, 21, BLUE, True, PP_ALIGN.CENTER)
    rect(slide, 0.62, 5.23, 12.05, 1.25, PANEL)
    text(slide, "Decisões-chave", 0.91, 5.56, 1.8, 0.28, 11, CYAN, True)
    text(slide, "VIN removido · data convertida em atributos temporais · alvo treinado em log1p · semente fixa · MMR investigado separadamente", 2.68, 5.5, 9.55, 0.44, 15, WHITE)
    footer(slide, 4)

    # 5 — MLP e espaço
    slide = base_slide(prs, "MLP e espaço de busca", "Modelo")
    rect(slide, 0.62, 1.52, 5.3, 4.92, PANEL)
    text(slide, "ARQUITETURA", 0.96, 1.87, 4.5, 0.3, 12, CYAN, True)
    architecture = [
        ("Entrada", "atributos processados", BLUE),
        ("Ocultas", "1 a 3 camadas densas", CYAN),
        ("Ativação", "ReLU ou tanh", TEAL),
        ("Saída", "Dense(1), linear", AMBER),
    ]
    for i, (a, b, c) in enumerate(architecture):
        y = 2.37 + i * 0.84
        rect(slide, 0.96, y, 4.55, 0.62, PANEL_2)
        rect(slide, 0.96, y, 0.08, 0.62, c, radius=False)
        text(slide, a, 1.22, y + 0.15, 1.15, 0.23, 12, WHITE, True)
        text(slide, b, 2.42, y + 0.15, 2.76, 0.24, 12, MUTED)
    text(slide, "Loss: MSE · seleção: MAE · Early Stopping", 0.96, 5.83, 4.55, 0.3, 12, WHITE, True)
    rect(slide, 6.2, 1.52, 6.47, 4.92, PANEL)
    text(slide, "HIPERPARÂMETROS", 6.56, 1.87, 5.4, 0.3, 12, CYAN, True)
    rich_lines(slide, [
        "Neurônios: 32, 64, 128 ou 256",
        "Learning rate: 10⁻⁴ a 10⁻²",
        "Batch size: 32, 64 ou 128",
        "Otimizador: Adam ou RMSprop",
        "Paciência: 8, 12, 16 ou 20",
        "Orçamento idêntico: 25 × 3 folds por estratégia",
    ], 6.56, 2.38, 5.45, 3.4, 16, 8)
    footer(slide, 5)

    # 6 — Buscas
    slide = base_slide(prs, "Random Search venceu; TPE foi mais consistente", "Comparação das buscas")
    add_bar_chart(slide, ["Random Search", "TPE"], [("Melhor MAE", (1027.26, 1072.11))], 0.63, 1.52, 6.15, 4.05, 1200)
    metric_card(slide, "Vantagem Random", "4,2%", 0.88, 5.79, 2.55, GREEN, "menor MAE")
    metric_card(slide, "Orçamento", "25 + 25", 3.68, 5.79, 2.55, BLUE, "3 folds cada")
    rect(slide, 7.08, 1.52, 5.59, 4.27, PANEL)
    text(slide, "MELHORES CONFIGURAÇÕES", 7.43, 1.87, 4.8, 0.28, 12, CYAN, True)
    text(slide, "Random Search", 7.43, 2.38, 2.1, 0.3, 16, WHITE, True)
    text(slide, "256 → 128 · tanh · RMSprop\nLR 0,0001804 · lote 128 · paciência 12", 7.43, 2.79, 4.5, 0.85, 14, MUTED)
    text(slide, "TPE", 7.43, 3.88, 2.1, 0.3, 16, WHITE, True)
    text(slide, "128 · ReLU · RMSprop\nLR 0,0001195 · lote 128 · paciência 8", 7.43, 4.29, 4.5, 0.85, 14, MUTED)
    rect(slide, 7.08, 6.02, 5.59, 0.52, BLUE)
    text(slide, "TPE concentrou mais tentativas em regiões boas.", 7.36, 6.15, 5.0, 0.24, 12, WHITE, True)
    footer(slide, 6)

    # 7 — Andrey (gráfico nativo)
    slide = base_slide(prs, "Andrey de Matos Gonçalves", "Análise individual de sensibilidade")
    text(slide, "Taxa de aprendizado (learning rate)", 0.64, 1.26, 5.5, 0.34, 15, CYAN, True)
    add_bar_chart(slide, ["0,0001", "0,0003", "0,001", "0,003", "0,01"],
                  [("MAE médio", (1033.38, 1035.03, 1140.12, 1357.22, 1872.51))],
                  0.63, 1.75, 7.85, 4.9, 2000)
    rect(slide, 8.72, 1.75, 3.95, 2.15, PANEL)
    text(slide, "MELHOR PONTO", 9.02, 2.06, 3.2, 0.25, 11, MUTED, True)
    text(slide, "1 × 10⁻⁴", 9.02, 2.43, 3.2, 0.48, 25, WHITE, True)
    text(slide, "MAE: US$ 1.033,38", 9.02, 3.05, 3.2, 0.27, 12, GREEN, True)
    rect(slide, 8.72, 4.11, 3.95, 2.54, PANEL)
    text(slide, "LEITURA", 9.02, 4.42, 3.2, 0.25, 11, MUTED, True)
    text(slide, "Maior impacto observado: amplitude de US$ 839. Taxas altas tornaram as atualizações instáveis.", 9.02, 4.82, 3.15, 1.36, 15, WHITE)
    footer(slide, 7)

    # 8-10 — demais sensibilidades
    sensitivity_slide(
        prs, 8, "Felipe Gabriel Souza Libório", "Tamanho do lote (batch size)",
        FIG_DIR / "sensibilidade_felipe_gabriel.png",
        "O efeito não foi monotônico. O lote 128 combinou menor erro e menor dispersão entre folds.",
        ("128", "MAE: US$ 1.027,26"),
    )
    sensitivity_slide(
        prs, 9, "Felipe Vieira Brazão e Silva", "Neurônios na primeira camada oculta",
        FIG_DIR / "sensibilidade_felipe_vieira.png",
        "256 foi a escolha parcimoniosa. Dobrar para 512 não trouxe ganho relevante de desempenho.",
        ("256 neurônios", "MAE: US$ 1.027,26"),
    )
    sensitivity_slide(
        prs, 10, "Gabriel Coelho Goes", "Paciência do Early Stopping",
        FIG_DIR / "sensibilidade_gabriel.png",
        "O limite de 80 épocas ocorreu antes da parada antecipada; por isso, o MAE permaneceu igual.",
        ("4 épocas", "Mesmo MAE · menor espera"),
    )

    # 11 — MMR
    slide = base_slide(prs, "A variável MMR concentra grande poder preditivo", "Ablação e baseline")
    rect(slide, 0.62, 1.52, 7.75, 4.92, WHITE)
    add_picture_contain(slide, FIG_DIR / "ablacao_mmr.png", 0.8, 1.7, 7.4, 4.56)
    rect(slide, 8.65, 1.52, 4.02, 4.92, PANEL)
    text(slide, "RESULTADO", 8.99, 1.87, 3.2, 0.27, 11, CYAN, True)
    text(slide, "+83,1%", 8.99, 2.35, 3.1, 0.62, 31, RED, True)
    text(slide, "no MAE ao retirar MMR", 8.99, 2.95, 3.15, 0.32, 13, MUTED)
    rich_lines(slide, [
        "Baseline MMR: US$ 1.085,97",
        "MLP com MMR: US$ 1.027,26",
        "MLP sem MMR: US$ 1.880,53",
    ], 8.99, 3.57, 3.1, 1.5, 14, 8)
    text(slide, "Uso legítimo somente se o MMR estiver disponível antes da venda.", 8.99, 5.42, 3.1, 0.65, 13, AMBER, True)
    footer(slide, 11)

    # 12 — Resultado final
    slide = base_slide(prs, "Desempenho no teste reservado", "Avaliação final")
    metric_card(slide, "MAE", "US$ 1.065,15", 0.63, 1.52, 3.65, BLUE, "erro absoluto médio")
    metric_card(slide, "RMSE", "US$ 1.652,57", 0.63, 3.03, 3.65, CYAN, "penaliza erros grandes")
    metric_card(slide, "R²", "0,9710", 0.63, 4.54, 3.65, GREEN, "97,1% da variância explicada")
    rect(slide, 4.62, 1.52, 8.05, 4.92, WHITE)
    add_picture_contain(slide, FIG_DIR / "curvas_modelo_final.png", 4.82, 1.72, 7.65, 4.52)
    text(slide, "75 épocas · melhor validação na época 63 · leve overfitting tardio controlado pelo Early Stopping", 4.8, 6.53, 7.72, 0.33, 11, MUTED, True, PP_ALIGN.CENTER)
    footer(slide, 12)

    # 13 — Respostas finais
    slide = base_slide(prs, "O que os experimentos responderam", "Síntese crítica")
    answers = [
        ("Busca", "Random Search encontrou o melhor ponto; TPE apresentou maior consistência.", BLUE),
        ("Impacto", "Learning rate foi o hiperparâmetro dominante e confirmou a importância do TPE.", CYAN),
        ("Generalização", "Não houve underfitting forte; houve apenas overfitting tardio e controlado.", TEAL),
        ("Mercado", "MMR é uma referência forte, mas a MLP ainda reduziu seu MAE em 5,4%.", AMBER),
    ]
    for i, (label, desc, color) in enumerate(answers):
        y = 1.5 + i * 1.23
        rect(slide, 0.62, y, 12.05, 0.95, PANEL)
        rect(slide, 0.62, y, 0.11, 0.95, color, radius=False)
        text(slide, label.upper(), 0.98, y + 0.21, 1.45, 0.26, 11, color, True)
        text(slide, desc, 2.48, y + 0.18, 9.68, 0.44, 16, WHITE, True)
    text(slide, "Limitações: uma única base, orçamento de 25 tentativas e dependência da disponibilidade pré-venda do MMR.", 0.72, 6.57, 11.8, 0.35, 12, MUTED, True, PP_ALIGN.CENTER)
    footer(slide, 13)

    # 14 — Encerramento
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    rect(slide, 0, 0, 0.18, 7.5, BLUE, radius=False)
    text(slide, "CONCLUSÃO", 0.75, 0.8, 3.0, 0.35, 12, CYAN, True)
    text(slide, "A MLP estimou preços com alta precisão, mas a análise crítica mostrou que o resultado depende tanto da otimização quanto da qualidade e disponibilidade dos atributos.", 0.75, 1.55, 11.65, 1.8, 30, WHITE, True)
    rect(slide, 0.75, 4.25, 11.82, 1.15, PANEL)
    text(slide, "Melhor modelo: Random Search · MAE final US$ 1.065,15 · R² 0,9710", 1.05, 4.61, 11.22, 0.38, 19, GREEN, True, PP_ALIGN.CENTER)
    text(slide, "Perguntas?", 0.75, 6.25, 4.0, 0.55, 26, WHITE, True)
    footer(slide, 14)

    prs.core_properties.title = "Regressão de preços de veículos com MLP"
    prs.core_properties.subject = "Random Search, TPE e análises individuais de sensibilidade"
    prs.core_properties.author = "Andrey, Felipe Gabriel, Felipe Vieira e Gabriel"
    prs.core_properties.keywords = "MLP, regressão, Random Search, TPE, veículos"
    prs.save(OUT_FILE)
    print(f"Apresentação criada: {OUT_FILE}")
    print(f"Slides: {len(prs.slides)}")


if __name__ == "__main__":
    main()
